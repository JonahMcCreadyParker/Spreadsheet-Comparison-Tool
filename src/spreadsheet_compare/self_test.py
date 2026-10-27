"""Unattended packaged-GUI verification; no user files or preferences are touched."""
import csv
import hashlib
import json
from pathlib import Path
import time
import traceback

from openpyxl import load_workbook, Workbook
from PySide6.QtCore import QTimer, Qt

from .core.models import MODES
from .gui.main_window import MainWindow


def run_self_test(app, output):
    folder = Path(output).resolve()
    folder.mkdir(parents=True, exist_ok=True)
    columns = ["Tag Name", "Address", "Data Type", "Comment", "Project"]
    original, revised = folder / "Demo Original.csv", folder / "Demo Revised.csv"
    for path, rows in ((original, [columns, ["TAG1", "S7:[PLC1]DB620,Real48,1", "REAL", "Old comment", "A"], ["OLD", "2", "LONG", "Removed", "A"]]),
                       (revised, [columns, ["NEW", "3", "LONG", "Added", "B"], ["TAG1", "S7:[PLC1]DB620,Real52,1", "LONG", "New comment", "B"]])):
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows(rows)
    hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (original, revised)}
    excel_inputs = []
    for path, formula in ((original, "=SUM(A1:A10)"), (revised, "=SUM(A1:A11)")):
        target = path.with_suffix(".xlsx")
        wb = Workbook()
        wb.active.title = "Tags"
        with path.open(encoding="utf-8-sig", newline="") as stream:
            for row in csv.reader(stream):
                wb.active.append(row)
        wb.create_sheet("Formula")["B1"] = formula
        wb.save(target)
        wb.close()
        excel_inputs.append(target)
        hashes[target.name] = hashlib.sha256(target.read_bytes()).hexdigest()
    window = MainWindow(persist=False)
    window.show()
    errors, observations = [], {}
    window.show_error = errors.append

    def select_key():
        for index in range(window.keys.count()):
            item = window.keys.item(index)
            if item.text() == "Tag Name":
                item.setCheckState(Qt.CheckState.Checked)
                return
        raise AssertionError("Tag Name key missing")

    def check_record():
        assert window.result.counts["modified"] == 1
        assert window.result.counts["added"] == window.result.counts["removed"] == 1
        observations[window.settings.mode] = dict(window.result.counts)

    def check_engineering():
        check_record()
        assert window.result.counts["fields"] == 3
        window.table.selectRow(0)

    def check_excel():
        wb = load_workbook(folder / "GUI Report.xlsx")
        assert "Differences" in wb.sheetnames
        assert wb["Differences"].max_row > 1
        wb.close()

    def check_excel_inputs():
        assert window.result.counts["sheets"] == 2
        assert any(d.category == "Formula" for d in window.result.differences)
        observations["Excel inputs / cell mode"] = dict(window.result.counts)

    def check_mixed():
        check_record()
        assert window.result.settings.sheet_pairs == [("CSV Data", "Tags")]
        observations["Mixed CSV / Excel records"] = dict(window.result.counts)

    def finish():
        assert {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (original, revised, *excel_inputs)} == hashes
        window.grab().save(str(folder / "Application.png"))
        (folder / "self-test.json").write_text(json.dumps({"passed": True, "modes": observations, "source_checksums_unchanged": True, "errors": errors}, indent=2), encoding="utf-8")
        window.close()
        app.exit(0)

    stages = [
        lambda: window.zones[0].accept_paths([str(original)]),
        lambda: window.zones[1].accept_paths([str(revised)]),
        lambda: window.mode.setCurrentText(MODES[1]), select_key,
        window.start_compare, check_record,
        lambda: window.export_to(folder / "GUI Report.xlsx"), check_excel,
        lambda: window.mode.setCurrentText(MODES[0]), window.start_compare,
        lambda: observations.update({MODES[0]: dict(window.result.counts)}),
        lambda: window.load_file(0, excel_inputs[0]), lambda: window.load_file(1, excel_inputs[1]),
        window.start_compare, check_excel_inputs,
        lambda: window.load_file(0, original),
        lambda: window.mode.setCurrentText(MODES[1]), select_key, window.start_compare, check_mixed,
        lambda: window.load_file(1, revised),
        lambda: window.mode.setCurrentText(MODES[2]), window.start_compare, check_engineering,
        lambda: window.export_to(folder / "GUI Report.csv"), finish,
    ]
    started = time.monotonic()
    timer = QTimer()
    timer.setInterval(50)

    def advance():
        try:
            if time.monotonic() - started > 60:
                raise TimeoutError("GUI self-test timed out")
            if errors:
                raise AssertionError("; ".join(errors))
            if window.worker or window._preview_pending:
                return
            if stages:
                stages.pop(0)()
        except Exception:
            timer.stop()
            (folder / "self-test.json").write_text(json.dumps({"passed": False, "error": traceback.format_exc(), "errors": errors}, indent=2), encoding="utf-8")
            if window.worker:
                window.worker.cancel()
                window.worker.wait(30000)
                window.task_finished()
            window.close()
            app.exit(1)

    timer.timeout.connect(advance)
    timer.start()
    return app.exec()
