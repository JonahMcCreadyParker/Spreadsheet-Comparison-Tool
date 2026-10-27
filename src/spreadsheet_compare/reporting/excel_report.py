import json
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .common import HEADERS, report_target, row_values, value_text
from ..core.models import InputError, checkpoint

COLORS = {"Added": "E2F0D9", "Removed": "FCE4D6", "Modified": "FFF2CC", "Unchanged": "E7E6E6", "Check": "FFF2CC", "Ambiguous": "FFF2CC"}


def export_excel(result, path, *, cancel=None, progress=None):
    wb = Workbook()
    wb.remove(wb.active)
    overflow = []

    def write_sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        for n, values in enumerate([headers], 1):
            ws.append(values)
        for index, values in enumerate(rows, 2):
            if index % 500 == 0:
                checkpoint(cancel)
                if progress:
                    progress(f"Writing {name}: {index - 1:,} rows…")
            if index > 1048576:
                raise InputError("Excel report exceeds the worksheet row limit. Export CSV instead.")
            for col, value in enumerate(values, 1):
                cell = ws.cell(index, col)
                if value is None:
                    continue
                if not isinstance(value, (int, float, bool)):
                    text = value_text(value)
                    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", lambda m: f"\\x{ord(m[0]):02x}", text)
                    if len(text) > 32767:
                        ref = f"{name}!{get_column_letter(col)}{index}"
                        for part, start in enumerate(range(0, len(text), 30000), 1):
                            overflow.append([ref, part, text[start:start + 30000]])
                        text = f"Long text: see Long Text sheet for {ref} ({len(text)} characters)."
                    cell.value = text
                    cell.data_type = "s"  # Imported formula expressions remain inert, literal strings.
                else:
                    cell.value = value
                if name != "Summary" and values and str(values[0]) in COLORS:
                    cell.fill = PatternFill("solid", fgColor=COLORS[str(values[0])])
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183B51")
        for col, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(col)].width = 54 if "Value" in str(header) or "Entry" in str(header) or header in ("Notes", "Value / interpretation") else 25
        return ws

    summary = [
        ("Application", "Spreadsheet Comparison Tool 1.0.0"),
        ("Original A", str(result.original.path)), ("Revised B", str(result.revised.path)),
        ("Comparison time", result.timestamp), ("Mode", result.settings.mode),
        ("Report scope", "All results, regardless of GUI filters"),
        ("Counting", "Unique records/cells; changed fields counts modified shared-record fields or cell attributes. Added/removed record payloads and structural changes are counted separately."),
        ("Blank display", "⟨blank cell⟩ = absent/None; ⟨empty text⟩ = literal empty string. Raw values are not normalised in this report."),
        ("Warnings / ambiguities", sum(d.severity != "Info" for d in result.warnings)),
        ("Notices", sum(d.severity == "Info" for d in result.warnings)),
        ("Structural changes", len(result.structural)),
        ("Interpretation", "Changes and engineering checks require review; they do not prove a defect. Blank/duplicate identities are excluded from confirmed record counts."),
        ("Formula safety", "All imported strings are literal text. No formulas or hyperlinks are created. Control characters are displayed as escaped text."),
        ("Import A", json.dumps(result.original.import_options, ensure_ascii=False)),
        ("Import B", json.dumps(result.revised.import_options, ensure_ascii=False)),
    ]
    summary += [(f"Count: {k}", v) for k, v in result.counts.items()]
    summary += [(f"Setting: {k}", json.dumps(v, ensure_ascii=False, default=str)) for k, v in result.settings.as_dict().items()]
    summary += [("Notice / warning", d.notes) for d in result.warnings]
    write_sheet("Summary", ["Property", "Value / interpretation"], summary)
    write_sheet("Differences", HEADERS, (row_values(d) for d in result.differences))
    if result.settings.mode != "General — Cell by Cell":
        for name, records in (("Added Records", result.added_records), ("Removed Records", result.removed_records)):
            # Each raw field gets its own column; avoid hiding records inside JSON blobs.
            fields = list(dict.fromkeys(k for r in records for k in r["Values"]))
            write_sheet(name, ["Worksheet", "Key", "Source row"] + fields,
                ([r["Worksheet"], r["Key"], r["Source row"]] + [value_text(r["Values"].get(h)) for h in fields] for r in records))
    write_sheet("Warnings", HEADERS, (row_values(d) for d in result.warnings))
    write_sheet("Structural Changes", HEADERS, (row_values(d) for d in result.structural))
    if overflow:
        write_sheet("Long Text", ["Source cell", "Part", "Literal text chunk"], list(overflow))
    checkpoint(cancel)
    with report_target(path, result, ".xlsx") as temp:
        wb.save(temp)
        checkpoint(cancel)
    wb.close()
    return str(path)
