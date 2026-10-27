from openpyxl.utils import get_column_letter

from .models import CellSnapshot, Difference, checkpoint
from .normalisation import same


def compare_cells(a, b, label, result, cancel=None):
    settings = result.settings
    contexts = []
    for sheet in (a, b):
        rows = {}
        for index, ((row, col), cell) in enumerate(sheet.cells.items()):
            if index % 1000 == 0:
                checkpoint(cancel)
            if cell.value is not None:
                rows.setdefault(row, {})[f"{get_column_letter(col)}{row}"] = cell.value
        contexts.append(rows)
    for index, position in enumerate(sorted(a.cells.keys() | b.cells.keys())):
        if index % 500 == 0:
            checkpoint(cancel)
        ca, cb = a.cells.get(position, CellSnapshot(style=a.default_style)), b.cells.get(position, CellSnapshot(style=b.default_style))
        row, col = position
        identity = f"{get_column_letter(col)}{row}"
        changes = []
        if settings.formulas or not (ca.formula or cb.formula):
            if not same(ca.value, cb.value, settings) or ca.formula != cb.formula:
                changes.append(("Formula" if ca.formula or cb.formula else "Value", "Value", ca.value, cb.value))
        if settings.formatting:
            for attribute in sorted(ca.style.keys() | cb.style.keys()):
                va, vb = ca.style.get(attribute, a.default_style.get(attribute)), cb.style.get(attribute, b.default_style.get(attribute))
                if va != vb:
                    changes.append(("Formatting", attribute, va, vb))
        if settings.hyperlinks and ca.hyperlink != cb.hyperlink:
            changes.append(("Hyperlink", "Hyperlink target", ca.hyperlink, cb.hyperlink))
        if not changes:
            if ca.value is not None or cb.value is not None:
                result.counts["unchanged"] += 1
            continue
        status = "Added" if ca.value is None and cb.value is not None else "Removed" if ca.value is not None and cb.value is None else "Modified"
        result.counts[status.lower()] += 1
        result.counts["fields"] += len(changes)
        for category, attribute, va, vb in changes:
            result.differences.append(Difference(status, label, identity, attribute, va, vb, category,
                notes="Same cell address; row insertions can cascade into multiple differences.",
                row_a=row if position in a.cells or row in contexts[0] else None, row_b=row if position in b.cells or row in contexts[1] else None,
                context_a=contexts[0].get(row, {}), context_b=contexts[1].get(row, {})))
    if a.merged != b.merged:
        result.structural.append(Difference("Modified", label, field="Merged ranges", original=", ".join(a.merged), revised=", ".join(b.merged), category="Worksheet structure"))
