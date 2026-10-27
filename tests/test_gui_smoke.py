import time

import pytest
from PySide6.QtCore import QMimeData, QPoint, QPointF, Qt, QUrl
from PySide6.QtGui import QDragEnterEvent, QDropEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog

from spreadsheet_compare.core.models import MODES
from spreadsheet_compare.gui.main_window import MainWindow
from spreadsheet_compare.gui.settings_dialog import SettingsDialog


def wait_idle(window, timeout=10000):
    until = time.monotonic() + timeout / 1000
    while window.worker or window._preview_pending:
        QApplication.processEvents()
        QTest.qWait(5)
        assert time.monotonic() < until, "GUI worker did not finish"
    QApplication.processEvents()


@pytest.fixture
def window(app, monkeypatch):
    w = MainWindow(persist=False)
    errors = []
    monkeypatch.setattr(w, "show_error", errors.append)
    w.test_errors = errors
    w.show()
    yield w
    wait_idle(w)
    w.close()


def drop_file(zone, path):
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(path))])
    enter = QDragEnterEvent(QPoint(10, 10), Qt.DropAction.CopyAction, mime, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(zone, enter)
    assert enter.isAccepted()
    event = QDropEvent(QPointF(10, 10), Qt.DropAction.CopyAction, mime, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(zone, event)
    assert event.isAccepted()


def test_gui_drop_browse_compare_filter_export(window, csv_book, tmp_path, monkeypatch):
    a = csv_book("a.csv", [["ID", "Value"], ["1", "a"], ["2", "b"]])
    b = csv_book("b.csv", [["ID", "Value"], ["2", "b"], ["1", "changed"]])
    assert not window.compare_button.isEnabled()
    drop_file(window.zones[0], a.path)
    wait_idle(window)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (str(b.path), ""))
    QTest.mouseClick(window.zones[1].browse, Qt.MouseButton.LeftButton)
    wait_idle(window)
    window.mode.setCurrentText(MODES[1])
    wait_idle(window)
    assert not window.compare_button.isEnabled()
    window.keys.item(0).setCheckState(Qt.CheckState.Checked)
    wait_idle(window)
    assert "0 rows with duplicate keys" in window.preview.toPlainText()
    assert window.compare_button.isEnabled()
    QTest.mouseClick(window.compare_button, Qt.MouseButton.LeftButton)
    wait_idle(window)
    assert window.result.counts["modified"] == 1
    assert window.model.rowCount() == 1
    window.table.selectRow(0)
    QApplication.processEvents()
    assert "changed" in window.context.toPlainText()
    window.search.setText("no such value")
    assert window.proxy.rowCount() == 0
    window.export_to(tmp_path / "report.xlsx")
    wait_idle(window)
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "report.xlsx")
    assert wb["Differences"].max_row == 2  # GUI filter does not restrict export.
    wb.close()
    window.export_to(tmp_path / "report.csv")
    wait_idle(window)
    assert (tmp_path / "report.csv").exists()
    assert not window.test_errors


def test_gui_engineering_proposes_key_and_swap(window, csv_book):
    a = csv_book("a.csv", [["Tag Name", "Address"], ["TAG1", "10"]])
    b = csv_book("b.csv", [["Tag Name", "Address"], ["TAG1", "20"], ["TAG2", "30"]])
    for side, book in enumerate((a, b)):
        window.load_file(side, book.path)
        wait_idle(window)
    window.mode.setCurrentText(MODES[2])
    wait_idle(window)
    assert window.settings.keys == ["Tag Name"]
    window.start_compare()
    wait_idle(window)
    assert window.result.counts["added"] == 1
    window.swap_files()
    wait_idle(window)
    window.start_compare()
    wait_idle(window)
    assert window.result.counts["removed"] == 1
    assert not window.test_errors


def test_drop_rejects_multiple_and_unsupported(window, tmp_path):
    zone = window.zones[0]
    messages = []
    zone.invalid.disconnect()
    zone.invalid.connect(messages.append)
    zone.accept_paths(["one.csv", "two.csv"])
    zone.accept_paths(["legacy.xls"])
    assert len(messages) == 2
    assert not window.books[0]


def test_settings_dialog_roundtrip(window):
    dialog = SettingsDialog(window.settings, ["ID", "Name"], ["Tag", "Label"], {}, window.books, window)
    dialog.mode.setCurrentText(MODES[1])
    dialog.mapping.cellWidget(0, 1).setCurrentText("Tag")
    dialog.mapping.cellWidget(1, 1).setCurrentText("Label")
    dialog.checks["trim"].setChecked(True)
    dialog.header_a.setValue(2)
    dialog.accept()
    assert dialog.settings.mappings == {"ID": "Tag", "Name": "Label"}
    assert dialog.settings.trim and dialog.settings.header_a == 2


def test_gui_large_import_keeps_event_loop_alive(window, csv_book):
    book = csv_book("large.csv", [["ID", "Value"]] + [[str(i), "x"] for i in range(12000)])
    from PySide6.QtCore import QTimer
    heartbeats = []
    timer = QTimer()
    timer.setInterval(1)
    timer.timeout.connect(lambda: heartbeats.append(1))
    timer.start()
    window.load_file(0, book.path)
    wait_idle(window)
    timer.stop()
    assert heartbeats
    assert window.books[0]
    assert not window.test_errors


def test_gui_mixed_excel_pairing(window, csv_book, excel_book):
    a = csv_book("a.csv", [["ID", "Value"], ["001", "1"]])
    b = excel_book("b.xlsx", {"Other": [["Other"]], "Table": [["ID", "Value"], ["001", 1]]})
    for side, book in enumerate((a, b)):
        window.load_file(side, book.path)
        wait_idle(window)
    assert window.scope.currentIndex() == 1
    window.sheet_b.setCurrentText("Table")
    wait_idle(window)
    window.mode.setCurrentText(MODES[1])
    wait_idle(window)
    window.keys.item(0).setCheckState(Qt.CheckState.Checked)
    wait_idle(window)
    window.start_compare()
    wait_idle(window)
    assert window.result.counts["modified"] == 1
    assert window.result.settings.sheet_pairs == [("CSV Data", "Table")]
    assert not window.test_errors


def test_gui_can_clear_engineering_proposed_key(window, csv_book):
    book = csv_book("a.csv", [["Tag Name", "Address"], ["TAG", "10"]])
    for side in (0, 1):
        window.load_file(side, book.path)
        wait_idle(window)
    window.mode.setCurrentText(MODES[2])
    wait_idle(window)
    window.keys.item(0).setCheckState(Qt.CheckState.Unchecked)
    wait_idle(window)
    assert not window.settings.keys
    assert not window.compare_button.isEnabled()


def test_gui_cancel_worker(window):
    from spreadsheet_compare.core.models import checkpoint
    def operation(cancel=None, progress=None):
        while not cancel():
            time.sleep(0.001)
        checkpoint(cancel)
    completed = []
    window.run_task(operation, completed.append)
    QTest.mouseClick(window.cancel_button, Qt.MouseButton.LeftButton)
    wait_idle(window)
    assert not completed
    assert "Cancelled" in window.status.text()
