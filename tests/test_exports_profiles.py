import csv
import hashlib
import json

import pytest
from openpyxl import load_workbook

from witthoft_compare.config.profiles import load_profiles, save_profiles
from witthoft_compare.core.compare_service import compare
from witthoft_compare.core.models import ComparisonSettings, Difference, InputError, Cancelled
from witthoft_compare.reporting.csv_report import export_csv
from witthoft_compare.reporting.excel_report import export_excel


def test_exports_inert_and_checksums(csv_book, records, tmp_path):
    payloads = ['=HYPERLINK("https://example.com","click")', '+CMD', '-CMD', '@SUM(1)', ' \t=1+1']
    a = csv_book("a.csv", [["ID", "Value"]] + [[str(i), "old"] for i in range(5)])
    b = csv_book("b.csv", [["ID", "Value"]] + [[str(i), p] for i, p in enumerate(payloads)])
    hashes = [hashlib.sha256(book.path.read_bytes()).hexdigest() for book in (a, b)]
    r = compare(a, b, records)
    xlsx, flat = tmp_path / "report.xlsx", tmp_path / "report.csv"
    export_excel(r, xlsx)
    export_csv(r, flat)
    wb = load_workbook(xlsx, data_only=False)
    assert wb.sheetnames == ["Summary", "Differences", "Added Records", "Removed Records", "Warnings", "Structural Changes"]
    assert [wb["Differences"].cell(i, 6).value for i in range(2, 7)] == payloads
    assert all(wb["Differences"].cell(i, 6).data_type == "s" for i in range(2, 7))
    assert all(cell.data_type != "f" and cell.hyperlink is None for ws in wb for row in ws for cell in row)
    assert wb["Differences"].freeze_panes == "A2"
    assert wb["Differences"].auto_filter.ref
    wb.close()
    assert flat.read_bytes().startswith(b'\xef\xbb\xbf')
    with flat.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    assert [row[5] for row in rows[1:]] == ["'" + p for p in payloads]
    assert [hashlib.sha256(book.path.read_bytes()).hexdigest() for book in (a, b)] == hashes


def test_cannot_overwrite_source(csv_book, records, tmp_path):
    a = csv_book("a.csv", [["ID"], ["1"]])
    r = compare(a, a, records)
    with pytest.raises(InputError, match="source"):
        export_csv(r, a.path)
    with pytest.raises(InputError, match="ending"):
        export_excel(r, tmp_path / "wrong.csv")


def test_export_cancel_is_atomic(csv_book, records, tmp_path):
    a = csv_book("a.csv", [["ID"], ["1"]])
    r = compare(a, a, records)
    path = tmp_path / "report.csv"
    path.write_text("original content")
    with pytest.raises(Cancelled):
        export_csv(r, path, cancel=lambda: True)
    assert path.read_text() == "original content"


def test_long_text_report(csv_book, records, tmp_path):
    a = csv_book("a.csv", [["ID", "v"], ["1", "a"]])
    b = csv_book("b.csv", [["ID", "v"], ["1", "=" + "x" * 40000]])
    r = compare(a, b, records)
    path = tmp_path / "report.xlsx"
    export_excel(r, path)
    wb = load_workbook(path)
    assert "Long Text" in wb.sheetnames
    chunks = {}
    for ref, part, text in list(wb["Long Text"].values)[1:]:
        chunks.setdefault(ref, []).append((part, text))
    restored = {ref: "".join(text for _, text in sorted(parts)) for ref, parts in chunks.items()}
    assert restored["Differences!F2"] == "=" + "x" * 40000
    assert json.loads(restored["Differences!M2"])["v"] == "=" + "x" * 40000
    wb.close()


def test_profile_recovery_and_round_trip(tmp_path):
    path = tmp_path / "settings.json"
    path.write_text("broken")
    assert load_profiles(path) == ({}, ComparisonSettings())
    path.write_text(json.dumps({"last": {"keys": [None, "ID"], "mode": "bad", "header_a": -20, "sheet_pairs": [1, ["x"]]}}))
    _, s = load_profiles(path)
    assert s.keys == ["ID"] and s.header_a == 1 and not s.sheet_pairs
    assert s.mode == ComparisonSettings().mode
    s.mappings = {"ID": "Tag"}
    save_profiles({"Custom": s}, s, path)
    profiles, last = load_profiles(path)
    assert profiles["Custom"] == last == s
